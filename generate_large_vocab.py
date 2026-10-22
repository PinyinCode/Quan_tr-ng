import json
import os
import time
from openai import OpenAI
import openpyxl

# Lấy biến môi trường từ GitHub Actions hoặc hệ thống
api_key = os.environ.get("DASHSCOPE_API_KEY")
base_url = os.environ.get(
    "QWEN_BASE_URL", "https://dashscope.aliyuncs.com/compatible-mode/v1"
)
model_name = os.environ.get("QWEN_MODEL", "qwen-plus")

client = OpenAI(api_key=api_key, base_url=base_url, timeout=60.0)

# Cấu hình lưu vào thư mục scripts/data_copy/
output_dir = "scripts/data_copy"
os.makedirs(output_dir, exist_ok=True)

# Lấy nhóm ngành từ workflow (mặc định 'Công Nghệ Cao & AI' nếu chạy cục bộ)
TARGET_INDUSTRY = os.environ.get("TARGET_INDUSTRY", "Công Nghệ Cao & AI")

# Xử lý tên file đẹp, giữ tiếng Việt có dấu (VD: "Công Nghệ Cao & AI" -> "Công_Nghệ_Cao_AI_Full.xlsx")
clean_name = (
    TARGET_INDUSTRY.strip()
    .replace("&", "")
    .replace("-", "_")
    .replace("  ", " ")
    .replace(" ", "_")
)
filename_prefix = f"{clean_name}_Full"

# Bản đồ cấu hình các chủ điểm/chủ đề đã được nhóm & tối ưu
INDUSTRY_TOPICS_MAP = {
    "Công Nghệ Cao & AI": [
        {
            "chu_diem": "Tri Tue Nhan Tao AI",
            "chu_de": (
                "Học máy, mạng neural, LLM và thị giác máy tính (机器学习与大模型)"
            ),
        },
        {
            "chu_diem": "Ban Dan Va Vi Mach",
            "chu_de": (
                "Chip bán dẫn, bo mạch, vi xử lý và linh kiện điện tử"
                " (半导体与芯片设计)"
            ),
        },
        {
            "chu_diem": "Tu Dong Hoa Va Robot",
            "chu_de": (
                "Cánh tay robot, cảm biến và điều khiển tự động hóa (机器人与自动化)"
            ),
        },
        {
            "chu_diem": "Ha Tang Dam May Va Bao Mat",
            "chu_de": (
                "Máy chủ, điện toán đám mây, API và an ninh mạng (云计算与网络安全)"
            ),
        },
    ],
    "Du Lịch - Khách Sạn & Xuất Nhập Cảnh": [
        {
            "chu_diem": "Mon An Gia Vi Va Nha Hang",
            "chu_de": (
                "Các món ăn Trung - Việt, gia vị, đặt bàn và phục vụ nhà hàng"
                " (菜肴、调料与餐饮)"
            ),
        },
        {
            "chu_diem": "Khach San Va Luu Tru",
            "chu_de": (
                "Đặt phòng, lễ tân, dọn phòng và tiện ích khách sạn"
                " (酒店预订与客房服务)"
            ),
        },
        {
            "chu_diem": "Thu Tuc Khai Bao Hai Quan",
            "chu_de": (
                "Visa, hộ chiếu, hải quan và an ninh sân bay (签证、海关与安检)"
            ),
        },
        {
            "chu_diem": "Huong Dan Vien Va Thang Canh",
            "chu_de": (
                "Giới thiệu danh lam thắng cảnh, di chuyển và văn hóa (景点与导游)"
            ),
        },
    ],
    "Giáo Dục & Du Học": [
        {
            "chu_diem": "Ho So Va Thu Tuc Du Hoc",
            "chu_de": (
                "Hồ sơ xin học bổng, visa du học và thủ tục nhập học (留学申请与签证)"
            ),
        },
        {
            "chu_diem": "Giang Day Va Hoc Tap",
            "chu_de": "Từ vựng lớp học, bài giảng và thảo luận nhóm (课堂教学与讨论)",
        },
        {
            "chu_diem": "Thi Kiem Tra Va Danh Gia",
            "chu_de": "Kỳ thi, chấm điểm, luận văn và bằng cấp (考试与论文)",
        },
        {
            "chu_diem": "Doi Song Sinh Vien Ky Xuc",
            "chu_de": (
                "Thuê nhà, ký túc xá, làm thêm và giao lưu văn hóa (住宿与打工)"
            ),
        },
    ],
    "Bán Hàng & Marketing": [
        {
            "chu_diem": "Tiep Can Khach Hang",
            "chu_de": (
                "Tìm kiếm khách hàng tiềm năng, chào hàng và quảng cáo (推销与广告)"
            ),
        },
        {
            "chu_diem": "Dam Phan Chot Don",
            "chu_de": "Thương lượng giá cả, chốt hợp đồng và ưu đãi (谈判与签单)",
        },
        {
            "chu_diem": "Cham Soc Khach Hang",
            "chu_de": (
                "Giải đáp thắc mắc, xử lý khiếu nại và hậu mãi (售后与客户服务)"
            ),
        },
        {
            "chu_diem": "Thi Truong Va Doanh So",
            "chu_de": (
                "Nghiên cứu thị trường, chiến dịch khuyến mãi và KPI sales"
                " (市场与销售额)"
            ),
        },
    ],
    "Văn Hóa & Ngày Lễ": [
        {
            "chu_diem": "Tet Nguyen Dan Va Xuan Ve",
            "chu_de": (
                "Tết cổ truyền, phong bao lì xì, chúc Tết và món ăn ngày Tết"
                " (春节与过年习俗)"
            ),
        },
        {
            "chu_diem": "Cac Ngay Le Truyen Thong",
            "chu_de": (
                "Tết Trung Thu, Tết Đoan Ngọ, Tết Nguyên Tiêu và Thanh Minh"
                " (传统节日与风俗)"
            ),
        },
        {
            "chu_diem": "Ngay Le Hien Dai Va Quoc Khanh",
            "chu_de": (
                "Ngày Quốc Khánh, Quốc tế Lao động, Phụ nữ và Nhà giáo"
                " (现代节日与法定假日)"
            ),
        },
        {
            "chu_diem": "Phong Tuc Va Giao Luu Van Hoa",
            "chu_de": (
                "Nghi lễ cưới hỏi, cúng bái, trà đạo và nét đẹp văn hóa"
                " (民俗文化与礼仪)"
            ),
        },
    ],
    "May Mặc": [
        {
            "chu_diem": "Kho nguyên phụ liệu",
            "chu_de": "Vải chính và vải lót (织物与里料)",
        },
        {
            "chu_diem": "Kho nguyên phụ liệu",
            "chu_de": "Chỉ may, khuy cúc và phụ liệu khác (辅料与线类)",
        },
        {
            "chu_diem": "Phòng cắt mẫu",
            "chu_de": "Trải vải, cắt vải và rập (铺布、裁剪与纸样)",
        },
        {
            "chu_diem": "Chuyền may sản xuất",
            "chu_de": "Công đoạn may ráp thân trước và thân sau (前后片缝制)",
        },
        {
            "chu_diem": "Chuyền may sản xuất",
            "chu_de": "Công đoạn ráp tay, cổ và lai áo (袖子、领子与下摆)",
        },
        {
            "chu_diem": "Quản lý chất lượng",
            "chu_de": "Kiểm tra lỗi thành phẩm KCS (成品质检与疵点)",
        },
        {
            "chu_diem": "Hoàn thiện và Đóng gói",
            "chu_de": "Là ủi, dập mác và đóng thùng (整烫、贴标与包装)",
        },
        {
            "chu_diem": "Quản lý xưởng sản xuất",
            "chu_de": "An toàn lao động, tiến độ và máy móc (安全、进度与设备)",
        },
    ],
    "Y Tế": [
        {
            "chu_diem": "Kham Va Dieutri",
            "chu_de": "Quy trình khám bệnh và hỏi triệu chứng (问诊与症状)",
        },
        {
            "chu_diem": "Duoc Pham",
            "chu_de": "Tên các loại thuốc và liều dùng (药品与用量)",
        },
        {
            "chu_diem": "Xet Nghiem Y Khoa",
            "chu_de": (
                "Xét nghiệm, chẩn đoán hình ảnh và xét nghiệm máu (检查与化验)"
            ),
        },
        {
            "chu_diem": "Cap Cuu Y Te",
            "chu_de": "Sơ cứu, xử lý sự cố khẩn cấp và phẫu thuật (急救与手术)",
        },
    ],
    "Nông - Lâm - Thủy Sản": [
        {
            "chu_diem": "Trong Trot Va Cay Trong",
            "chu_de": (
                "Các loại cây trồng, giống cây, kỹ thuật gieo trồng và chăm sóc"
                " (农作物与种植)"
            ),
        },
        {
            "chu_diem": "Phan Bon Va Thuoc Tru Sau",
            "chu_de": (
                "Phân bón hóa học/hữu cơ, thuốc bảo vệ thực vật và cải tạo đất"
                " (化肥与农药)"
            ),
        },
        {
            "chu_diem": "Chan Nuoi Va Thuy San",
            "chu_de": (
                "Gia súc, gia cầm, nuôi trồng thủy hải sản và thức ăn chăn nuôi"
                " (畜牧与水产养殖)"
            ),
        },
        {
            "chu_diem": "Nong Nghiep Cong Nghe Cao",
            "chu_de": (
                "Hệ thống tưới tự động, nhà kính và máy móc nông nghiệp"
                " (现代农业与机械)"
            ),
        },
    ],
    "Cơ Khí & Chế Tạo": [
        {
            "chu_diem": "Gia Cong Co Khi Va Che Tao",
            "chu_de": (
                "Cắt kim loại, hàn, đúc, gia công chi tiết và vật liệu công"
                " nghiệp (机械加工与制造)"
            ),
        },
        {
            "chu_diem": "Thiet Bi Va May Moc",
            "chu_de": (
                "Dây chuyền lắp ráp, động cơ, máy nén khí và bảo trì thiết bị"
                " (工业设备与维护)"
            ),
        },
        {
            "chu_diem": "An Toan Lao Dong Industrial",
            "chu_de": (
                "Trang bị bảo hộ, phòng chống cháy nổ và quy định an toàn nhà"
                " xưởng (工业安全与防护)"
            ),
        },
        {
            "chu_diem": "Quan Ly San Xuat Va Khung Hai",
            "chu_de": (
                "Định mức sản xuất, quản lý chất lượng và vận hành nhà máy"
                " (工厂生产管理)"
            ),
        },
    ],
    "Hàng Không Vũ Trụ": [
        {
            "chu_diem": "Kieu Dang Va Cau Tao May Bay",
            "chu_de": (
                "Động cơ phản lực, cánh máy bay, buồng lái và hệ thống phi hành"
                " (飞机结构与发动机)"
            ),
        },
        {
            "chu_diem": "Kiem Soat Khong Luu Va An Ninh",
            "chu_de": (
                "Tháp điều khiển, đài radar, tuyến bay và an toàn hàng không"
                " (空管与航线管制)"
            ),
        },
        {
            "chu_diem": "Cong Nghe Vu Tru Va Ve Tinh",
            "chu_de": (
                "Tên lửa đẩy, vệ tinh nhân tạo, trạm vũ trụ và quỹ đạo (卫星与航天器)"
            ),
        },
        {
            "chu_diem": "Bao Duong Va Ky Thuat Hang Khong",
            "chu_de": (
                "Bảo dưỡng máy bay, kiểm tra định kỳ MRO và linh kiện thay thế"
                " (航空维修与检测)"
            ),
        },
    ],
}

default_topics = [
    {
        "chu_diem": f"Nghiep Vu {TARGET_INDUSTRY}",
        "chu_de": f"Thuật ngữ chuyên môn chính ngành {TARGET_INDUSTRY}",
    },
    {
        "chu_diem": f"Giao Tiep Cong Viec {TARGET_INDUSTRY}",
        "chu_de": f"Các tình huống giao tiếp thực tế ngành {TARGET_INDUSTRY}",
    },
    {
        "chu_diem": f"Quy Trinh Va Quan Ly",
        "chu_de": (
            f"Quy trình làm việc và giải quyết sự cố ngành {TARGET_INDUSTRY}"
        ),
    },
    {
        "chu_diem": f"Thiet Bi Va Cong Cu",
        "chu_de": (
            f"Trang thiết bị và ứng dụng thực tế trong ngành {TARGET_INDUSTRY}"
        ),
    },
]

industry_topics = INDUSTRY_TOPICS_MAP.get(TARGET_INDUSTRY, default_topics)


def generate_vocabulary_batch(chu_diem, chu_de):
  prompt = f"""Bạn là một chuyên gia ngôn ngữ và giảng viên tiếng Trung chuyên nghiệp.
Hãy tạo danh sách khoảng 50 từ vựng và thuật ngữ tiếng Trung thực tế thuộc lĩnh vực: {TARGET_INDUSTRY}
- Chủ điểm: {chu_diem}
- Chủ đề: {chu_de}

Yêu cầu BẮT BUỘC:
1. Từ vựng chuẩn xác, có tính ứng dụng cao thực tế trong đời sống, công việc.
2. Trả về đúng định dạng JSON là một mảng (array) chứa các object:
[
  {{
    "hsk": "HSK 3",
    "vi": "Nghĩa tiếng Việt",
    "zh": "Từ tiếng Trung",
    "py": "Pinyin đầy đủ có dấu"
  }},
  ...
]
"""
  max_retries = 3
  for attempt in range(max_retries):
    try:
      response = client.chat.completions.create(
          model=model_name,
          messages=[{"role": "user", "content": prompt}],
          temperature=0.8,
          response_format={"type": "json_object"},
      )
      data = json.loads(response.choices[0].message.content)
      if isinstance(data, dict):
        for k, v in data.items():
          if isinstance(v, list):
            data = v
            break
      return data
    except Exception as e:
      print(f"⚠️ Lỗi tạo từ vựng chủ đề {chu_de} (lần {attempt + 1}): {e}")
      time.sleep(2)
  return []


def generate_example_sentence(vi, zh):
  prompt = f"""Bạn là chuyên gia ngôn ngữ giỏi tiếng Trung.
Hãy đặt 01 câu giao tiếp tự nhiên, câu dài hoàn chỉnh có chứa từ vựng '{zh}' (nghĩa là: '{vi}').

Trả về đúng định dạng JSON:
{{
  "zh": "Câu tiếng Trung dài và tự nhiên",
  "py": "Pinyin đầy đủ có dấu của cả câu",
  "vi": "Bản dịch tiếng Việt hoàn chỉnh của câu"
}}
"""
  try:
    response = client.chat.completions.create(
        model=model_name,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.8,
        response_format={"type": "json_object"},
    )
    data = json.loads(response.choices[0].message.content)
    return data["vi"], data["zh"], data["py"]
  except Exception:
    return (
        f"Chúng ta cần chú ý xử lý vấn đề {vi} này một cách cẩn thận.",
        f"我们需要仔细处理这个{zh}问题。",
        f"Wǒmen xūyào zǐxì chǔlǐ zhè ge {zh} wèntí.",
    )


file_output = os.path.join(output_dir, f"{filename_prefix}.xlsx")
wb = openpyxl.Workbook()
sheet = wb.active
sheet.title = "Vocabulary"

headers = [
    "STT",
    "HSK",
    "Chủ điểm",
    "Chủ đề",
    "Tiếng Việt",
    "Tiếng Trung",
    "Pinyin",
    "Luyện tập",
    "Check đúng/sai",
]
for col_num, header_title in enumerate(headers, 1):
  sheet.cell(row=1, column=col_num, value=header_title)

current_row = 2
global_stt = 1

print(
    f"🚀 Bắt đầu tạo bộ từ vựng lớn cho nhóm ngành: [{TARGET_INDUSTRY}] (Tên"
    f" file: {filename_prefix}.xlsx)..."
)

for topic in industry_topics:
  c_diem = topic["chu_diem"]
  c_de = topic["chu_de"]

  print(
      f"\n📦 Đang sáng tạo từ vựng cho chủ điểm: [{c_diem}] - Chủ đề:"
      f" [{c_de}]..."
  )
  vocab_list = generate_vocabulary_batch(c_diem, c_de)
  print(f"   -> Đã tạo thành công {len(vocab_list)} từ vựng.")

  for item in vocab_list:
    hsk = item.get("hsk", "HSK 3")
    vi = item.get("vi")
    zh = item.get("zh")
    py = item.get("py")

    if not zh or not vi:
      continue

    # 1. HÀNG 1: TỪ VỰNG GỐC
    sheet.cell(row=current_row, column=1, value=global_stt)
    sheet.cell(row=current_row, column=2, value=hsk)
    sheet.cell(row=current_row, column=3, value=c_diem)
    sheet.cell(row=current_row, column=4, value=c_de)
    sheet.cell(row=current_row, column=5, value=vi)
    sheet.cell(row=current_row, column=6, value=zh)
    sheet.cell(row=current_row, column=7, value=py)
    current_row += 1
    global_stt += 1

    # 2. HÀNG 2: CÂU DÀI CHỨA TỪ ĐÓ (Dành cho từ ngắn <= 12 ký tự)
    if len(str(zh)) <= 12:
      ex_vi, ex_zh, ex_py = generate_example_sentence(vi, zh)

      sheet.cell(row=current_row, column=1, value=global_stt)
      sheet.cell(row=current_row, column=2, value=hsk)
      sheet.cell(row=current_row, column=3, value=c_diem)
      sheet.cell(row=current_row, column=4, value=c_de)
      sheet.cell(row=current_row, column=5, value=ex_vi)
      sheet.cell(row=current_row, column=6, value=ex_zh)
      sheet.cell(row=current_row, column=7, value=ex_py)
      current_row += 1
      global_stt += 1

      time.sleep(0.4)

wb.save(file_output)
print(
    f"\n🎉 HOÀN TẤT! Đã xuất file thành công tại thư mục copy: {file_output}"
    f" (Tổng cộng {current_row - 1} dòng dữ liệu)."
)
